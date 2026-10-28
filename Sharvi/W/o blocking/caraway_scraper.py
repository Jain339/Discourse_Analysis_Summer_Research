#!/usr/bin/env python3
"""
Caraway & Jain -- Phase 2: data collection (per-country version)
================================================================

Collects official sustainability-related communications published between
1 Jan 2024 and 31 May 2026 from the websites in the classification sheet,
ONE COUNTRY AT A TIME, then combines the countries into one month-sorted workbook.

    python caraway_scraper.py --list                     # show countries and sources
    python caraway_scraper.py --country USA --probe      # health-check one country's sites
    python caraway_scraper.py --country USA              # scrape one country
    python caraway_scraper.py --combine                  # merge every finished country

In a notebook:  from caraway_scraper import main;  main(["--country", "USA"])

Each country run writes  outputs/countries/<Country>/
    articles.csv       one row per kept article (link, date, month, keywords...)
    corpus.jsonl       full text + metadata (input for Step 2 preprocessing)
    source_log.csv     what each site produced, and how many pages were BLOCKED
    <Country>_Articles.xlsx
--combine writes  outputs/Caraway_Jain_Articles_ALL.xlsx  (+ articles_all.csv, corpus_all.jsonl)

The classification sheet can be the .xlsx OR two CSV exports of its tabs
(a file with "General" in its name and one with "Codewords" in its name).

Getting past blocking (details in README.md)
--------------------------------------------
* Run on your own computer, not Codespaces: many .gov sites refuse cloud servers.
* --js [--browser chrome]: render JavaScript sites with Playwright (a real browser).
* Wayback fallback (on by default): if a page is blocked, its Internet Archive
  copy is used instead; blocked sites can also be DISCOVERED through the archive.
* Official APIs where they exist (US Federal Register, GOV.UK).
* manual_urls/<Country>.csv: links you collected by hand (or with Claude in
  Chrome) are fetched, filtered and dated like everything else.
"""
from __future__ import annotations

import argparse
import datetime as dt
import gzip
import io
import json
import logging
import re
import sqlite3
import threading
import time
import zlib
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, replace
from pathlib import Path
from urllib import robotparser
from urllib.parse import parse_qs, parse_qsl, quote, urlencode, urljoin, urlparse, urlunparse

import pandas as pd
import requests
import trafilatura
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("caraway")

# ─────────────────────────────────────────────────────────────────────────────
# 1. SETTINGS  (edit these)
# ─────────────────────────────────────────────────────────────────────────────
DEFAULT_SHEET = "Caraway_Jain_Data_Classification_Sheet.xlsx"
DEFAULT_START = "2024-01-01"
DEFAULT_END = "2026-05-31"
CONTACT_EMAIL = "your.email@mail.utoronto.ca"  # shown to site admins in the User-Agent

# Analysis periods (comment NJ1). Start inclusive, end exclusive.
PERIODS = [
    ("Pre-withdrawal", "2024-01-01", "2025-01-01"),
    ("Transition", "2025-01-01", "2026-01-01"),
    ("Post-withdrawal", "2026-01-01", "2026-06-01"),
]

# Global North / South grouping as in the Step 1 text, plus continent (sheet note).
COUNTRY_INFO = {
    "USA": ("Global North", "North America"),
    "Canada": ("Global North", "North America"),
    "UK": ("Global North", "Europe"),
    "EU": ("Global North", "Europe"),
    "New Zealand": ("Global North", "Oceania"),
    "South Korea": ("Global North", "Asia"),
    "China": ("Global South", "Asia"),
    "India": ("Global South", "Asia"),
    "Singapore": ("Global South", "Asia"),
    "Hong Kong": ("Global South", "Asia"),
}

BROAD_CATEGORY_HINT = "adjective"   # the broad indicator-word column of the Codewords tab
MIN_BROAD_TERMS_FOR_STRONG = 3

SOURCE_TYPE_LABELS = {"gov": "Govt press release", "ministry": "Ministry of Environment", "news": "News"}

WAYBACK = "https://web.archive.org"
FEDREG_API = "https://www.federalregister.gov/api/v1/documents.json"

# Per-site settings, matched by URL prefix against the URLs in the sheet.
# Fields are those of the Source dataclass (section 2).
SOURCE_OVERRIDES: dict[str, dict] = {
    # ---------------- USA (all three sites block cloud servers) ----------------
    "https://www.state.gov/": dict(
        strategy="listing", page_url_template="{base}page/{n}/", page_start=2,
        wayback_prefix="www.state.gov/releases/",
        link_pattern=r"state\.gov/(releases/[^?#]+/20\d\d/\d\d/[^/?#]+|[a-z0-9-]{25,})/?$",
        notes="Live listing first; Internet Archive if blocked."),
    "https://2021-2025.state.gov/": dict(
        strategy="listing", page_url_template="{base}page/{n}/", page_start=2,
        link_pattern=r"2021-2025\.state\.gov/[a-z0-9-]{25,}/?$",
        notes="Archived State Dept site for pre-2025 releases."),
    "https://www.epa.gov/newsreleases": dict(
        strategy="listing", page_param="page", page_start=0, wayback_prefix="www.epa.gov/newsreleases/",
        link_pattern=r"epa\.gov/newsreleases/[a-z0-9-]{12,}/?$",
        notes="Live listing first; Internet Archive if blocked."),
    # ---------------- Canada ----------------
    "https://www.pm.gc.ca/en/news": dict(
        strategy="listing", page_param="page", page_start=0, wayback_prefix="www.pm.gc.ca/en/news/",
        link_pattern=r"/en/news/[^/]+/20\d\d/\d\d/\d\d/", notes="Drupal ?page=N"),
    "https://www.canada.ca/en/news/advanced-news-search": dict(
        strategy="listing", page_param="idx", page_start=0, page_step=10,
        fixed_params={"start": "{start}", "end": "{end}"},
        link_pattern=r"/news/20\d\d/\d\d/", notes="date-filtered search, idx offset"),
    # ---------------- UK ----------------
    "https://www.gov.uk/search/": dict(strategy="govuk_api", notes="GOV.UK Search API"),
    # ---------------- EU ----------------
    "https://environment.ec.europa.eu/news": dict(
        strategy="listing", page_param="page", page_start=0,
        link_pattern=r"environment\.ec\.europa\.eu/news/[^/?#]+-20\d\d-\d\d-\d\d_en$",
        notes="?page=N is zero-based; date is in the URL slug"),
    "https://ec.europa.eu/commission/presscorner": dict(
        render_js=True, link_pattern=r"presscorner/detail/en/(ip|statement|speech|qanda)_\d\d_\d+",
        notes="JavaScript app: run with --js on your own computer."),
    # ---------------- China ----------------
    "https://english.mee.gov.cn/": dict(
        strategy="listing", page_url_template="{base}index_{n}.shtml", page_start=1,
        link_pattern=r"/20\d{4}/(t\d{8}_\d+\.s?html?|[PU]\d+\.pdf)$",
        notes="Few items; many releases are PDFs (now extracted)."),
    "https://english.www.gov.cn/news": dict(link_pattern=r"/news/20\d{4}/\d\d/content_"),
    # ---------------- New Zealand ----------------
    "https://www.beehive.govt.nz/releases": dict(
        strategy="listing", page_param="page", page_start=0, wayback_prefix="www.beehive.govt.nz/release/",
        link_pattern=r"beehive\.govt\.nz/release/[^/?#]+$", notes="?page=N zero-based, ~10 per page"),
    "https://www.dpmc.govt.nz/": dict(wayback_prefix="www.dpmc.govt.nz/news/"),
    # ---------------- Singapore ----------------
    "https://www.sgpc.gov.sg": dict(render_js=True, notes="JavaScript-rendered: run with --js"),
    "https://www.mse.gov.sg/latest-news": dict(
        link_pattern=r"mse\.gov\.sg/latest-news/[^/?#]{12,}/?$", notes="pagination guessed automatically"),
    # ---------------- Hong Kong ----------------
    "https://www.info.gov.hk/gia/general/": dict(
        strategy="hk_daily", title_prefilter=True, link_pattern=r"/gia/general/\d{6}/\d{2}/P\d+\.htm",
        notes="Daily index of ALL HK government releases; only keyword-matching titles are fetched."),
    "https://www.cmab.gov.hk/": dict(
        allowed_domains=["info.gov.hk"], link_pattern=r"info\.gov\.hk/gia/general/\d{6}/\d{2}/P\d+\.htm",
        notes="Links point to info.gov.hk (duplicates of the daily index are removed at combine)."),
    "https://www.eeb.gov.hk/": dict(
        allowed_domains=["info.gov.hk"], link_pattern=r"info\.gov\.hk/gia/general/\d{6}/\d{2}/P\d+\.htm"),
    "https://www.epd.gov.hk/": dict(allowed_domains=["info.gov.hk"]),
    # ---------------- India ----------------
    "https://www.pib.gov.in/": dict(
        enabled=False, notes="ASP.NET forms: collect links into manual_urls/India.csv (see README)."),
    "https://moef.gov.in/": dict(wayback_prefix="moef.gov.in/"),
    # ---------------- South Korea ----------------
    "https://mcee.go.kr/": dict(link_pattern=r"mcee\.go\.kr/.*(boardId|seq|articleId|nttId)=\d+",
                                notes="Board pages; file-download links are skipped"),
}

# Sources not in the sheet that Step 1 needs ("executive orders").
EXTRA_SOURCES: list[dict] = [
    dict(country="USA", source_type="Govt press release", url="https://2021-2025.state.gov/press-releases/",
         notes="State Dept pages from before 20 Jan 2025 (not in the sheet)."),
    dict(country="USA", source_type="Executive order / Presidential document",
         url=FEDREG_API, strategy="federal_register",
         notes="Official Federal Register API: executive orders, memoranda, proclamations."),
]

FEDREG_DOC_TYPES = ["executive_order", "memorandum", "proclamation", "determination", "notice"]

# ─────────────────────────────────────────────────────────────────────────────
# 2. DATA CLASSES
# ─────────────────────────────────────────────────────────────────────────────
@dataclass
class Source:
    country: str
    source_type: str
    url: str
    strategy: str = "auto"            # auto|listing|sitemap|wayback|govuk_api|hk_daily|federal_register|manual
    link_pattern: str | None = None   # regex an article URL must match
    page_param: str | None = None
    page_start: int = 0
    page_step: int = 1
    page_url_template: str | None = None
    fixed_params: dict = field(default_factory=dict)
    allowed_domains: list = field(default_factory=list)  # extra domains article links may live on
    wayback_prefix: str | None = None  # where articles live, for Internet Archive discovery
    title_prefilter: bool = False      # only fetch links whose title matches a codeword
    slug_prefilter: bool = False       # only fetch URLs whose slug matches a codeword
    render_js: bool = False
    enabled: bool = True
    notes: str = ""
    manual_rows: list = field(default_factory=list, repr=False)

    @property
    def site(self) -> str:
        return urlparse(self.url).netloc.lower().removeprefix("www.")


@dataclass
class Candidate:
    url: str
    title_hint: str = ""
    date: dt.date | None = None
    precision: str = ""      # day | month
    date_source: str = ""
    fetch_url: str | None = None  # where to get the text if different from url
    doc_type: str = ""


@dataclass
class Page:
    body: bytes
    ctype: str
    via: str             # live | js | wayback | cache:<via>

    @property
    def text(self) -> str:
        return self.body.decode("utf-8", "replace")

    @property
    def is_pdf(self) -> bool:
        return "pdf" in self.ctype or self.body[:5] == b"%PDF-"


# ─────────────────────────────────────────────────────────────────────────────
# 3. READ THE CLASSIFICATION SHEET (.xlsx or two .csv files)
# ─────────────────────────────────────────────────────────────────────────────
URL_RE = re.compile(r"https?://[^\s,;]+")


def find_sheet(path: str | None) -> Path:
    """Return the .xlsx, or the folder holding the General/Codewords CSVs."""
    if path and path not in (".", "./"):
        p = Path(path)
        if p.exists():
            return p
        raise FileNotFoundError(f"{p.resolve()} not found")
    cwd = Path(".")
    if (cwd / DEFAULT_SHEET).exists():
        return cwd / DEFAULT_SHEET
    xlsx = [f for f in cwd.glob("*.xlsx") if "classification" in f.name.lower()]
    if xlsx:
        return xlsx[0]
    return cwd


def read_sheet(src: Path, tab: str) -> pd.DataFrame:
    if src.is_file() and src.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(src, sheet_name=tab, dtype=str)
    folder = src if src.is_dir() else src.parent
    matches = sorted(f for f in folder.glob("*.csv") if tab.lower() in f.name.lower())
    if not matches:
        raise FileNotFoundError(
            f"Classification sheet not found. Put {DEFAULT_SHEET} in {folder.resolve()}, or two CSVs "
            f"with 'General' and 'Codewords' in their names. CSVs there now: {[f.name for f in folder.glob('*.csv')]}")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return pd.read_csv(matches[0], dtype=str, encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode {matches[0]}")


def _find_col(columns, *prefixes):
    for c in columns:
        if any(str(c).strip().lower().startswith(p) for p in prefixes):
            return c
    return None


def apply_overrides(src: Source) -> Source:
    best = max((p for p in SOURCE_OVERRIDES if src.url.startswith(p)), key=len, default=None)
    return replace(src, **SOURCE_OVERRIDES[best]) if best else src


def load_sources(sheet: Path, manual_dir: Path = Path("manual_urls")) -> list[Source]:
    df = read_sheet(sheet, "General")
    country_col = _find_col(df.columns, "countries")
    gov_col = _find_col(df.columns, "govt press", "government press")
    min_col = _find_col(df.columns, "ministry")
    news_cols = [c for c in df.columns if str(c).strip().lower().startswith("news")]
    sources, seen = [], set()

    def add(country, stype, cell):
        if not isinstance(cell, str):
            return
        for url in URL_RE.findall(cell):
            url = url.rstrip(".,)")
            if url not in seen:
                seen.add(url)
                sources.append(apply_overrides(Source(country=country, source_type=stype, url=url)))

    for _, row in df.iterrows():
        country = row.get(country_col)
        if isinstance(country, str) and country.strip() and not country.strip().lower().startswith(("countries", "regulatory", "news")):
            add(country.strip(), SOURCE_TYPE_LABELS["gov"], row.get(gov_col))
            add(country.strip(), SOURCE_TYPE_LABELS["ministry"], row.get(min_col))
        for nc in news_cols:
            add("International (news)", SOURCE_TYPE_LABELS["news"], row.get(nc))
    for extra in EXTRA_SOURCES:
        if extra["url"] not in seen:
            sources.append(apply_overrides(Source(**extra)))
    sources += load_manual_sources(manual_dir)
    return sources


def load_manual_sources(manual_dir: Path) -> list[Source]:
    """manual_urls/<Country>.csv with column 'url' (optional: date, title, source_type)."""
    out = []
    if not manual_dir.exists():
        return out
    for f in sorted(manual_dir.glob("*.csv")):
        if f.name.startswith("_"):  # _TEMPLATE.csv etc.
            continue
        df = pd.read_csv(f, dtype=str).fillna("")
        if "url" not in df.columns:
            log.warning("%s has no 'url' column, skipped", f)
            continue
        country = f.stem.replace("_", " ")
        for stype, grp in df.groupby(df["source_type"] if "source_type" in df.columns else pd.Series("Manual list", index=df.index)):
            out.append(Source(country=country, source_type=stype or "Manual list", url=f"manual://{f.name}#{stype}",
                              strategy="manual", manual_rows=grp.to_dict("records"),
                              notes=f"{len(grp)} links from {f.name}"))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# 4. KEYWORDS
# ─────────────────────────────────────────────────────────────────────────────
@dataclass
class Keyword:
    term: str
    category: str
    country: str | None
    patterns: list


TITLES = {"dame", "sir", "dr", "dr.", "mr", "mr.", "ms", "ms.", "mrs", "hon", "hon."}


def _pattern(term: str) -> re.Pattern:
    if re.fullmatch(r"[A-Z0-9&]{2,8}", term):  # acronyms: case-sensitive; COP also matches COP28
        return re.compile(rf"(?<![A-Za-z]){re.escape(term)}\d*(?![A-Za-z])")
    body = r"[\s\-]+".join(re.escape(w) for w in term.split())
    return re.compile(rf"\b{body}(?:s|es)?\b", re.IGNORECASE)


def name_variants(name: str) -> list[str]:
    """'Dame Angela Eagle' -> + 'Angela Eagle'; 'Michael S. Regan' -> + 'Michael Regan';
    'Grace FU Hai Yien' -> + 'Grace Fu'."""
    words = [w for w in name.split() if w.lower() not in TITLES]
    variants = {name, " ".join(words)}
    no_initials = [w for w in words if not re.fullmatch(r"[A-Z]\.?", w)]
    variants.add(" ".join(no_initials))
    if len(no_initials) >= 4:
        variants.add(" ".join(no_initials[:2]))
    return [v for v in variants if len(v.split()) >= 2]


def load_keywords(sheet: Path) -> list[Keyword]:
    df = read_sheet(sheet, "Codewords")
    cats = list(df.columns[:4])
    ccol = df.columns[4] if len(df.columns) > 4 else None
    if ccol is not None:
        df[ccol] = df[ccol].ffill()
    out, seen = [], set()
    for _, row in df.iterrows():
        for cat in cats:
            cell = row.get(cat)
            if not isinstance(cell, str) or not cell.strip():
                continue
            is_person = cat == cats[3]
            country = row.get(ccol) if (ccol is not None and is_person) else None
            m = re.match(r"^(.*?)\s*\(([^)]+)\)\s*$", cell.strip())
            for t in ([m.group(1), m.group(2)] if m else [cell.strip()]):
                if t and (t.lower(), cat) not in seen:
                    seen.add((t.lower(), cat))
                    terms = name_variants(t) if is_person else [t]
                    out.append(Keyword(t, str(cat), country if isinstance(country, str) else None,
                                       [_pattern(x) for x in terms]))
    return out


def match_keywords(text: str, keywords: list[Keyword]) -> dict:
    by_cat: dict[str, set] = defaultdict(set)
    countries = set()
    for kw in keywords:
        if any(p.search(text) for p in kw.patterns):
            by_cat[kw.category].add(kw.term)
            if kw.country:
                countries.add(kw.country)
    strong = False
    for c, v in by_cat.items():
        if BROAD_CATEGORY_HINT not in c.lower():
            strong = strong or bool(v)
        else:
            specific = {t for t in v if not any(t != o and t.lower() in o.lower() for o in v)}
            strong = strong or len(specific) >= MIN_BROAD_TERMS_FOR_STRONG
    return {"by_cat": by_cat, "distinct": sum(len(v) for v in by_cat.values()),
            "strong": strong, "person_countries": sorted(countries)}


# ─────────────────────────────────────────────────────────────────────────────
# 5. DATES, PERIODS, DOCUMENT TYPES
# ─────────────────────────────────────────────────────────────────────────────
def _mkdate(y, m, d=1):
    try:
        out = dt.date(int(y), int(m), int(d))
        return out if 2000 <= out.year <= 2100 else None
    except (ValueError, TypeError):
        return None


def iso_date(s) -> dt.date | None:
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(s or ""))
    return _mkdate(*m.groups()) if m else None


def date_from_url(url: str):
    path = urlparse(url).path
    for rx in (r"(?<!\d)(20\d{2})[/-](\d{2})[/-](\d{2})(?!\d)",
               r"(?<!\d)(20\d{2})(\d{2})/(\d{2})(?!\d)",
               r"(?<![\dA-Za-z])[tPU]?0?(20\d{2})(\d{2})(\d{2})(?=[_\d])"):
        m = re.search(rx, path)
        if m and (d := _mkdate(*m.groups())):
            return d, "day"
    m = re.search(r"(?<!\d)(20\d{2})/(\d{2})(?!\d)", path)
    if m and (d := _mkdate(*m.groups())):
        return d, "month"
    for seg in path.split("/"):
        m = re.fullmatch(r"(20\d{2})(\d{2})", seg)
        if m and (d := _mkdate(*m.groups())):
            return d, "month"
    return None, ""


_MONTHS = ("jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec|january|february|march|april|june|"
           "july|august|september|october|november|december")
_TEXT_DATES = [
    (re.compile(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})"), "ymd"),
    (re.compile(r"(?<!\d)(\d{1,2})/(\d{1,2})/(20\d{2})"), "dmy_num"),
    (re.compile(rf"(\d{{1,2}})(?:st|nd|rd|th)?\s+({_MONTHS})\.?,?\s+(20\d{{2}})", re.I), "dmy"),
    (re.compile(rf"({_MONTHS})\.?\s+(\d{{1,2}}),?\s+(20\d{{2}})", re.I), "mdy"),
]


def date_from_text(s: str) -> dt.date | None:
    for rx, kind in _TEXT_DATES:
        m = rx.search(s)
        if not m:
            continue
        try:
            if kind == "ymd":
                return _mkdate(*m.groups())
            if kind == "dmy_num":  # 23/10/2026 (HK, NZ, UK, India style)
                d, mo, y = m.groups()
                return _mkdate(y, mo, d)
            d, mon, y = m.groups() if kind == "dmy" else (m.group(2), m.group(1), m.group(3))
            return dt.datetime.strptime(f"{d} {mon[:3]} {y}", "%d %b %Y").date()
        except ValueError:
            continue
    return None


def month_end(d: dt.date) -> dt.date:
    return dt.date(d.year + (d.month == 12), d.month % 12 + 1, 1) - dt.timedelta(days=1)


def outside_range(d, precision, start, end) -> bool:
    if d is None:
        return False
    if precision == "month":
        return month_end(d) < start or d.replace(day=1) > end
    return d < start or d > end


def period_of(d) -> str:
    if not d:
        return ""
    for label, a, b in PERIODS:
        if dt.date.fromisoformat(a) <= d < dt.date.fromisoformat(b):
            return label
    return ""


DOC_TYPE_RULES = [
    ("Executive order", r"\bexecutive order\b"),
    ("Speech / remarks", r"\b(speech|remarks|address|keynote|opening statement|toast|lecture)\b"),
    ("Interview / doorstop", r"\b(interview|doorstop|transcript|press conference|q\s?&\s?a|qanda)\b"),
    ("Joint / ministerial statement", r"\b(statement|joint declaration|communiqu[eé]|declaration)\b"),
    ("Policy / regulatory", r"\b(regulation|regulatory|rule|bill|act|directive|consultation|policy|strategy|"
                            r"roadmap|framework|plan|standard|guidance|legislation|proposal)\b"),
]


def classify_doc_type(title: str, url: str, source_type: str, hint: str = "") -> str:
    if hint:
        return hint
    hay = f"{title} {urlparse(url).path.replace('-', ' ').replace('_', ' ')}"
    for label, rx in DOC_TYPE_RULES:
        if re.search(rx, hay, re.I):
            return label
    return "News article" if source_type == SOURCE_TYPE_LABELS["news"] else "Press release"


# ─────────────────────────────────────────────────────────────────────────────
# 6. POLITE, CACHED FETCHING  (+ block detection, Wayback fallback, real browser)
# ─────────────────────────────────────────────────────────────────────────────
TRACKING = re.compile(r"^(utm_|fbclid|gclid|mc_|fontsize$)", re.I)
BLOCK_MARKERS = re.compile(
    r"just a moment\.\.\.|cf-chl|attention required! \| cloudflare|<title>\s*access denied|request rejected|"
    r"incapsula incident|_incapsula_resource|pardon our interruption|are you a robot|verify you are human|"
    r"please enable (javascript and )?cookies|bot detection|captcha-delivery|perimeterx|ddos protection by",
    re.I)
DAY = 24 * 3600


def normalize_url(url: str) -> str:
    p = urlparse(url)
    q = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True) if not TRACKING.match(k)]
    return urlunparse((p.scheme, p.netloc.lower(), p.path, "", urlencode(q), ""))


def looks_blocked(status: int, body: bytes) -> bool:
    if status in (401, 403, 406, 429, 202):
        return True
    head = body[:8000].decode("utf-8", "replace")
    return len(body) < 60000 and bool(BLOCK_MARKERS.search(head))


class Fetcher:
    """requests + robots.txt + per-domain delay + SQLite cache + Wayback fallback + optional browser."""

    def __init__(self, cache_path, delay, offline=False, use_js=False, browser="chromium",
                 headed=False, wayback=True, email=CONTACT_EMAIL):
        self.delay, self.offline, self.use_js = delay, offline, use_js
        self.browser_channel, self.headed, self.wayback = browser, headed, wayback
        self.ua = f"CarawayJainResearch/0.2 (academic research; contact: {email})"
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": self.ua, "Accept-Language": "en"})
        Path(cache_path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(cache_path, check_same_thread=False)
        self.db.execute("CREATE TABLE IF NOT EXISTS pages2 (url TEXT PRIMARY KEY, fetched REAL, ctype TEXT, via TEXT, body BLOB)")
        self.lock = threading.Lock()
        self.domain_locks = defaultdict(threading.Lock)
        self.last_hit: dict[str, float] = {}
        self.robots: dict[str, robotparser.RobotFileParser | None] = {}
        self.robots_text: dict[str, str] = {}
        self.stats = Counter()
        self.last_reason = threading.local()
        self._js = threading.local()
        self._browsers = []
        self.blocked_hosts = Counter()

    # cache ---------------------------------------------------------------
    def _cache_get(self, url, max_age):
        with self.lock:
            row = self.db.execute("SELECT fetched, ctype, via, body FROM pages2 WHERE url=?", (url,)).fetchone()
        if row and (max_age is None or self.offline or time.time() - row[0] < max_age):
            return Page(zlib.decompress(row[3]), row[1] or "", "cache:" + (row[2] or ""))
        return None

    def _cache_put(self, url, page: Page):
        with self.lock:
            self.db.execute("INSERT OR REPLACE INTO pages2 VALUES (?,?,?,?,?)",
                            (url, time.time(), page.ctype, page.via, zlib.compress(page.body)))
            self.db.commit()

    # robots.txt ------------------------------------------------------------
    def allowed(self, url) -> bool:
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        if base not in self.robots:
            rp = None
            res = self._http(base + "/robots.txt", respect_robots=False)
            if isinstance(res, Page) and "<html" not in res.text[:500].lower():
                rp = robotparser.RobotFileParser()
                rp.parse(res.text.splitlines())
                self.robots_text[base] = res.text
            self.robots[base] = rp
        rp = self.robots[base]
        return True if rp is None else rp.can_fetch(self.ua, url)

    def sitemaps_for(self, url):
        p = urlparse(url)
        self.allowed(url)
        return re.findall(r"(?im)^\s*sitemap:\s*(\S+)", self.robots_text.get(f"{p.scheme}://{p.netloc}", ""))

    # network ---------------------------------------------------------------
    def _throttle(self, netloc, delay=None):
        with self.domain_locks[netloc]:
            wait = self.last_hit.get(netloc, 0) + (self.delay if delay is None else delay) - time.time()
            if wait > 0:
                time.sleep(wait)
            self.last_hit[netloc] = time.time()

    def _http(self, url, respect_robots=True, via="live"):
        """Return Page, or a reason string: robots | blocked | http_<code> | failed."""
        if respect_robots and not self.allowed(url):
            return "robots"
        netloc = urlparse(url).netloc
        reason = "failed"
        tries = 5 if netloc.endswith("archive.org") else 3
        for attempt in range(tries):
            self._throttle(netloc, self.delay * 2 if netloc.endswith("archive.org") else None)
            try:
                r = self.session.get(url, timeout=150 if netloc.endswith("archive.org") else 40)
            except requests.RequestException as e:
                log.debug("error %s: %s", url, e)
                time.sleep(2 ** attempt)
                continue
            body = r.content
            if body[:2] == b"\x1f\x8b":
                try:
                    body = gzip.decompress(body)
                except OSError:
                    pass
            if r.status_code == 200 and not looks_blocked(200, body):
                return Page(body, r.headers.get("content-type", "").lower(), via)
            if r.status_code in (202, 429, 500, 502, 503, 504):
                reason = "blocked" if r.status_code in (202, 429) else f"http_{r.status_code}"
                time.sleep((20 if netloc.endswith("archive.org") else 6) * (attempt + 1))
                continue
            return "blocked" if looks_blocked(r.status_code, body) else f"http_{r.status_code}"
        return reason

    def _render_js(self, url):
        if not hasattr(self._js, "browser"):
            from playwright.sync_api import sync_playwright  # optional dependency
            pw = sync_playwright().start()
            kw = {"headless": not self.headed}
            if self.browser_channel in ("chrome", "msedge"):
                kw["channel"] = self.browser_channel  # your installed Chrome / Edge
            self._js.pw, self._js.browser = pw, pw.chromium.launch(**kw)
            self._browsers.append((pw, self._js.browser))
        if not self.allowed(url):
            return "robots"
        self._throttle(urlparse(url).netloc, max(self.delay, 3))
        page = self._js.browser.new_page()
        try:
            resp = page.goto(url, wait_until="networkidle", timeout=60000)
            html = page.content().encode("utf-8")
            status = resp.status if resp else 200
            if looks_blocked(status, html):
                page.wait_for_timeout(8000)  # some checks clear themselves after a few seconds
                html = page.content().encode("utf-8")
                if looks_blocked(200, html):
                    return "blocked"
            return Page(html, "text/html", "js")
        except Exception as e:
            log.debug("js render failed %s: %s", url, e)
            return "failed"
        finally:
            page.close()

    def _wayback(self, url, when: dt.date | None):
        stamp = (when or dt.date(2025, 6, 1)).strftime("%Y%m%d")
        res = self._http(f"{WAYBACK}/web/{stamp}id_/{url}", respect_robots=False, via="wayback")
        return res

    def get(self, url, max_age=None, render_js=False, when=None, allow_wayback=True) -> Page | None:
        url = normalize_url(url)
        cached = self._cache_get(url, max_age)
        if cached is not None:
            self.stats["cache_hits"] += 1
            self.last_reason.value = ""
            return cached
        if self.offline:
            self.last_reason.value = "offline"
            return None
        if render_js and not self.use_js:
            res = "needs_js"
        else:
            res = self._render_js(url) if render_js else self._http(url)
            if not isinstance(res, Page) and self.use_js and not render_js and res == "blocked":
                res = self._render_js(url)  # a real browser often passes where requests is refused
        if not isinstance(res, Page) and res in ("blocked", "failed", "needs_js", "http_403") \
                and self.wayback and allow_wayback:
            self.blocked_hosts[urlparse(url).netloc] += 1
            wb = self._wayback(url, when)
            res = wb if isinstance(wb, Page) else res
        if isinstance(res, Page):
            self.stats[f"fetched_{res.via}"] += 1
            self._cache_put(url, res)  # only good pages are cached; blocked ones are retried next run
            self.last_reason.value = ""
            return res
        self.stats[res] += 1
        self.last_reason.value = res
        return None

    def close(self):
        for pw, b in self._browsers:
            try:
                b.close(); pw.stop()
            except Exception:
                pass
        self.db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 7. DISCOVERY STRATEGIES (find candidate article URLs)
# ─────────────────────────────────────────────────────────────────────────────
SKIP_PATH = re.compile(r"/(search|login|signin|contact|about|privacy|terms|accessibility|sitemap|subscribe|"
                       r"rss|feed|tag|tags|category|categories|author|topics?)(/|$)", re.I)
SKIP_EXT = re.compile(r"\.(jpe?g|png|gif|svg|zip|docx?|xlsx?|pptx?|mp4|mp3|xml|json|css|js)$", re.I)
SKIP_ANY = re.compile(r"download|readDownloadFile|attach|fileDown", re.I)
NEXT_TEXT = re.compile(r"^\s*(next|next page|older|older posts|›|»|>|next ›|next »)\s*$", re.I)
CC_SUFFIXES = (".gov.uk", ".gc.ca", ".gov.cn", ".gov.hk", ".govt.nz", ".gov.sg", ".go.kr", ".gov.in",
               ".europa.eu", ".co.uk", ".com.cn")


def same_site(url, src: Source) -> bool:
    host = urlparse(url).netloc.lower().removeprefix("www.")
    base = src.site
    n = 3 if base.endswith(CC_SUFFIXES) else 2
    root = ".".join(base.split(".")[-n:])
    if host == base or host == root or host.endswith("." + root):
        return True
    return any(host == d or host.endswith("." + d) for d in src.allowed_domains)


def looks_like_article(url, src: Source, listing_url: str) -> bool:
    if SKIP_ANY.search(url) and not url.lower().endswith(".pdf"):
        return False
    if src.link_pattern:
        return re.search(src.link_pattern, url) is not None
    p = urlparse(url)
    if normalize_url(url) == normalize_url(listing_url) or SKIP_PATH.search(p.path) or SKIP_EXT.search(p.path):
        return False
    if re.search(r"[?&](page|idx|start|p|pageIndex)=", url) or re.search(r"/(page|index)[-_/]?\d+(\.\w+)?/?$", p.path, re.I):
        return False
    segs = [s for s in p.path.split("/") if s]
    return len(segs) >= 2 and (len(segs[-1]) >= 12 or bool(re.search(r"\d", segs[-1])))


def extract_links(html, page_url, src: Source) -> list[Candidate]:
    soup = BeautifulSoup(html, "lxml")
    out, seen = [], set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href.startswith(("#", "mailto:", "javascript:", "tel:")):
            continue
        rel = a.get("rel") or []
        if "next" in rel or "prev" in rel or NEXT_TEXT.match(a.get_text(" ", strip=True) or "-"):
            continue
        url = normalize_url(urljoin(page_url, href))
        if url in seen or not same_site(url, src) or not looks_like_article(url, src, page_url):
            continue
        seen.add(url)
        title = " ".join(a.get_text(" ", strip=True).split())
        d, prec = date_from_url(url)
        dsrc = "url" if d else ""
        if not d:
            node = a
            for _ in range(3):
                node = node.parent
                if node is None:
                    break
                txt = node.get_text(" ", strip=True)
                if len(txt) > 400:
                    break
                if (d := date_from_text(txt)):
                    prec, dsrc = "day", "listing page"
                    break
        out.append(Candidate(url, title, d, prec, dsrc))
    return out


def find_next_link(html, page_url):
    soup = BeautifulSoup(html, "lxml")
    tag = soup.find("a", rel=lambda v: v and "next" in v) or soup.find("link", rel=lambda v: v and "next" in v)
    if tag is None:
        for a in soup.find_all("a", href=True):
            label = a.get_text(" ", strip=True) or a.get("aria-label", "") or a.get("title", "")
            if NEXT_TEXT.match(label) or re.search(r"\bnext\b", a.get("aria-label", ""), re.I):
                tag = a
                break
    if tag is not None and tag.get("href") and not tag["href"].startswith(("#", "javascript:")):
        return normalize_url(urljoin(page_url, tag["href"]))
    return None


def _with_params(url, params):
    p = urlparse(url)
    q = dict(parse_qsl(p.query, keep_blank_values=True))
    q.update(params)
    return urlunparse((p.scheme, p.netloc, p.path, "", urlencode(q), ""))


def page_url(src: Source, i, start, end):
    fixed = {k: v.format(start=start.isoformat(), end=end.isoformat()) for k, v in src.fixed_params.items()}
    base = _with_params(src.url, fixed) if fixed else src.url
    if src.page_url_template:
        if i == 0:
            return base
        folder = base if base.endswith("/") else base.rsplit("/", 1)[0] + "/"
        return src.page_url_template.format(base=folder, n=src.page_start + (i - 1) * src.page_step)
    if src.page_param:
        return _with_params(base, {src.page_param: str(src.page_start + i * src.page_step)})
    return base


PAGINATION_GUESSES = [  # (page_param, page_start) or template; tried when a page has no "next" link
    dict(page_param="page", page_start=0), dict(page_param="page", page_start=1),
    dict(page_param="pageIndex", page_start=1), dict(page_param="p", page_start=1),
    dict(page_url_template="{base}page/{n}/", page_start=2),
    dict(page_url_template="{base}index_{n}.shtml", page_start=1),
]


def guess_pagination(src, fetcher, first_links, start, end):
    known = {c.url for c in first_links}
    for g in PAGINATION_GUESSES:
        trial = replace(src, **g)
        url = page_url(trial, 1, start, end)
        if normalize_url(url) == normalize_url(page_url(src, 0, start, end)):
            continue
        page = fetcher.get(url, max_age=DAY, render_js=src.render_js, allow_wayback=False)
        if page and not page.is_pdf:
            new = [c for c in extract_links(page.text, url, src) if c.url not in known]
            if len(new) >= 3:
                log.info("  %s: pagination detected (%s)", src.site, g)
                return trial
    return None


def discover_listing(src, fetcher, start, end, max_pages):
    seen = set()
    url = page_url(src, 0, start, end)
    paginated = bool(src.page_param or src.page_url_template)
    for i in range(max_pages):
        page = fetcher.get(url, max_age=DAY, render_js=src.render_js, allow_wayback=False)
        if not page:
            if i == 0 and src.wayback_prefix and fetcher.wayback:
                log.info("  %s listing blocked (%s); switching to Internet Archive discovery",
                         src.site, fetcher.last_reason.value)
                yield from discover_wayback(src, fetcher, start, end, max_pages)
            break
        cands = [c for c in extract_links(page.text, url, src) if c.url not in seen]
        if not cands:
            break
        for c in cands:
            seen.add(c.url)
            yield c
        dated = [c for c in cands if c.date]
        if dated and all(outside_range(c.date, c.precision, start, dt.date.max) for c in dated):
            break  # whole page older than the start date
        if paginated:
            url = page_url(src, i + 1, start, end)
            continue
        nxt = find_next_link(page.text, url)
        if not nxt and i == 0:
            guessed = guess_pagination(src, fetcher, cands, start, end)
            if guessed:
                src, paginated = guessed, True
                nxt = page_url(src, 1, start, end)
        if not nxt or nxt in seen:
            break
        url = nxt


def _parse_sitemap(xml):
    soup = BeautifulSoup(xml, "xml")
    kids = [(s.loc.text.strip(), s.lastmod.text.strip() if s.lastmod else "") for s in soup.find_all("sitemap") if s.loc]
    urls = [(u.loc.text.strip(), u.lastmod.text.strip() if u.lastmod else "") for u in soup.find_all("url") if u.loc]
    return kids, urls


def discover_sitemap(src, fetcher, start, end, max_sitemaps=300):
    p = urlparse(src.url)
    base = f"{p.scheme}://{p.netloc}"
    queue = fetcher.sitemaps_for(src.url) or [base + "/sitemap.xml", base + "/sitemap_index.xml"]
    seen = set()
    while queue and len(seen) < max_sitemaps:
        sm = queue.pop(0)
        if sm in seen:
            continue
        seen.add(sm)
        page = fetcher.get(sm, max_age=DAY, allow_wayback=False)
        if not page or "<" not in page.text[:200]:
            continue
        kids, urls = _parse_sitemap(page.text)
        queue += [loc for loc, lm in kids if not iso_date(lm) or iso_date(lm) >= start]
        for loc, lm in urls:
            if not same_site(loc, src) or (src.link_pattern and not re.search(src.link_pattern, loc)):
                continue
            if iso_date(lm) and iso_date(lm) < start:
                continue
            d, prec = date_from_url(loc)
            yield Candidate(normalize_url(loc), "", d, prec, "url" if d else "")


def discover_auto(src, fetcher, start, end, max_pages):
    got = False
    if src.link_pattern:
        for c in discover_sitemap(src, fetcher, start, end):
            got = True
            yield c
    if not got:
        yield from discover_listing(src, fetcher, start, end, max_pages)


def discover_wayback(src, fetcher, start, end, max_pages, cdx_limit=1500):
    """List every archived URL under wayback_prefix captured in the period (Internet Archive CDX API).
    Works for sites that block us, because we never contact the site itself."""
    prefix = src.wayback_prefix or (urlparse(src.url).netloc + urlparse(src.url).path.rsplit("/", 1)[0] + "/")
    resume, rows_seen = None, 0
    while True:
        params = {"url": prefix, "matchType": "prefix", "output": "json", "fl": "timestamp,original",
                  "filter": "statuscode:200", "collapse": "urlkey", "limit": str(cdx_limit), "showResumeKey": "true",
                  "from": start.strftime("%Y%m%d"),
                  "to": (end + dt.timedelta(days=45)).strftime("%Y%m%d")}
        if resume:
            params["resumeKey"] = resume
        url = f"{WAYBACK}/cdx/search/cdx?" + urlencode(params)
        page = fetcher.get(url, max_age=7 * DAY, allow_wayback=False)
        if not page:
            log.warning("  Internet Archive is not responding for %s (%s). It is often overloaded: "
                        "try again later; pages already listed are cached.", prefix, fetcher.last_reason.value)
            return
        try:
            rows = json.loads(page.text or "[]")
        except json.JSONDecodeError:
            return
        resume = None
        if rows and rows[0] == ["timestamp", "original"]:
            rows = rows[1:]
        if len(rows) >= 2 and rows[-2] == []:  # resume key is the last row after an empty row
            resume = rows[-1][0]
            rows = rows[:-2]
        for row in rows:
            if len(row) < 2:
                continue
            ts, orig = row[0], row[1]
            orig = normalize_url(orig if orig.startswith("http") else "https://" + orig)
            if src.link_pattern and not re.search(src.link_pattern, orig):
                continue
            if SKIP_EXT.search(urlparse(orig).path):
                continue
            slug = urlparse(orig).path.rstrip("/").rsplit("/", 1)[-1].replace("-", " ").replace("_", " ")
            d, prec = date_from_url(orig)
            first_seen = _mkdate(ts[:4], ts[4:6], ts[6:8])
            rows_seen += 1
            hint = slug if src.slug_prefilter else ""
            if d:
                yield Candidate(orig, hint, d, prec, "url")
            else:  # no date in the URL: first capture is an upper bound; the page's own date is used if found
                yield Candidate(orig, hint, first_seen, "", "first archive capture")
        if not resume:
            break
    log.info("  %s: %d archived URLs listed", prefix, rows_seen)


def discover_govuk(src, fetcher, start, end, max_pages):
    qs = parse_qs(urlparse(src.url).query)
    orgs = qs.get("organisations[]") or qs.get("organisations") or []
    for offset in range(0, max_pages * 100, 100):
        params = {"filter_content_purpose_supergroup": "news_and_communications",
                  "filter_public_timestamp": f"from:{start},to:{end}", "order": "-public_timestamp",
                  "count": "100", "start": str(offset), "fields": "title,link,public_timestamp,format"}
        url = "https://www.gov.uk/api/search.json?" + urlencode(params) + "".join(f"&filter_organisations={o}" for o in orgs)
        page = fetcher.get(url, max_age=DAY, allow_wayback=False)
        if not page:
            break
        results = json.loads(page.text).get("results", [])
        for r in results:
            fmt = (r.get("format") or "").replace("_", " ")
            hint = {"speech": "Speech / remarks", "written statement": "Joint / ministerial statement",
                    "oral statement": "Joint / ministerial statement"}.get(fmt, "")
            yield Candidate(normalize_url(urljoin("https://www.gov.uk", r["link"])), r.get("title", ""),
                            iso_date(r.get("public_timestamp")), "day", "GOV.UK API", doc_type=hint)
        if len(results) < 100:
            break


def discover_hk_daily(src, fetcher, start, end, max_pages):
    p = urlparse(src.url)
    base = f"{p.scheme}://{p.netloc}/gia/general/"
    d = start
    while d <= end:
        idx = f"{base}{d:%Y%m}/{d:%d}.htm"
        page = fetcher.get(idx, max_age=None if (dt.date.today() - d).days > 2 else DAY)
        if page:
            for c in extract_links(page.text, idx, src):
                c.date, c.precision, c.date_source = d, "day", "daily index"
                yield c
        d += dt.timedelta(days=1)


FEDREG_FIELDS = ("title", "html_url", "raw_text_url", "signing_date", "publication_date", "document_number")


def discover_federal_register(src, fetcher, start, end, max_pages):
    """US presidential documents (executive orders etc.) from the official Federal Register API.
    One query per document type, so the type is known without relying on extra fields."""
    for doc_type in FEDREG_DOC_TYPES:
        label = "Executive order" if doc_type == "executive_order" else f"Presidential {doc_type.replace('_', ' ')}"
        base = [("per_page", "100"), ("order", "oldest"), ("conditions[type][]", "PRESDOCU"),
                ("conditions[presidential_document_type]", doc_type),
                ("conditions[publication_date][gte]", start.isoformat()),
                ("conditions[publication_date][lte]", (end + dt.timedelta(days=21)).isoformat())]
        url = src.url + "?" + urlencode(base + [("fields[]", f) for f in FEDREG_FIELDS])
        page = fetcher.get(url, max_age=DAY, allow_wayback=False)
        if not page and fetcher.last_reason.value == "http_400":  # API rejected a field: use its defaults
            log.info("  Federal Register rejected the field list; retrying with default fields")
            url = src.url + "?" + urlencode(base)
            page = fetcher.get(url, max_age=DAY, allow_wayback=False)
        for _ in range(max_pages):
            if not page:
                log.warning("  Federal Register query failed for %s (%s)", doc_type, fetcher.last_reason.value)
                break
            data = json.loads(page.text)
            for r in data.get("results", []):
                d = iso_date(r.get("signing_date")) or iso_date(r.get("publication_date"))
                yield Candidate(r["html_url"], r.get("title", ""), d, "day",
                                "signing date" if r.get("signing_date") else "publication date",
                                fetch_url=r.get("raw_text_url") or r["html_url"], doc_type=label)
            nxt = data.get("next_page_url")
            if not nxt:
                break
            page = fetcher.get(nxt, max_age=DAY, allow_wayback=False)


def discover_manual(src, fetcher, start, end, max_pages):
    for row in src.manual_rows:
        u = (row.get("url") or "").strip()
        if not u.startswith("http"):
            continue
        d = iso_date(row.get("date")) or date_from_text(row.get("date", "") or "")
        yield Candidate(normalize_url(u), row.get("title", ""), d, "day" if d else "",
                        "manual list" if d else "")


STRATEGIES = {"auto": discover_auto, "listing": discover_listing, "govuk_api": discover_govuk,
              "hk_daily": discover_hk_daily, "wayback": discover_wayback,
              "federal_register": discover_federal_register, "manual": discover_manual,
              "sitemap": lambda s, f, a, b, m: discover_sitemap(s, f, a, b)}


# ─────────────────────────────────────────────────────────────────────────────
# 8. PROCESS ONE SOURCE
# ─────────────────────────────────────────────────────────────────────────────
def pdf_text(body: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(body))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:40])
    except Exception as e:
        log.debug("pdf failed: %s", e)
        return ""


def extract(page: Page, url: str):
    if page.is_pdf:
        text = pdf_text(page.body)
        first = next((l.strip() for l in text.splitlines() if len(l.strip()) > 15), "")
        return {"title": first[:200], "text": text, "date": None}
    html = page.text
    if "text/plain" in page.ctype or (not re.search(r"<(html|body|div|p)\b", html[:3000], re.I)):
        return {"title": "", "text": html, "date": None}
    raw = trafilatura.extract(html, url=url, output_format="json", with_metadata=True,
                              include_comments=False, include_tables=True)
    meta = json.loads(raw) if raw else {}
    if not meta.get("text"):  # e.g. Federal Register raw text wrapped in <pre>
        meta["text"] = BeautifulSoup(html, "lxml").get_text("\n", strip=True)
    return meta


def resolve_date(cand: Candidate, meta: dict):
    url_d, url_prec = date_from_url(cand.url)
    if url_d and url_prec == "day":
        return url_d, "url"
    if cand.date and cand.date_source in ("GOV.UK API", "daily index", "signing date", "publication date", "manual list"):
        return cand.date, cand.date_source
    if (md := iso_date(meta.get("date"))):
        return md, "page metadata"
    if cand.date and cand.precision == "day":
        return cand.date, cand.date_source
    if url_d:
        return url_d, "url (month only)"
    if cand.date:
        return cand.date, cand.date_source  # first archive capture (approximate)
    return None, ""


def run_source(src: Source, fetcher: Fetcher, keywords, start, end, max_pages, limit):
    stats = Counter()
    records = []
    discover = STRATEGIES.get(src.strategy, discover_auto)
    seen = set()
    try:
        for cand in discover(src, fetcher, start, end, max_pages):
            if cand.url in seen:
                continue
            seen.add(cand.url)
            stats["candidates"] += 1
            if outside_range(cand.date, cand.precision, start, end) and cand.date_source != "first archive capture":
                stats["skipped_out_of_range"] += 1
                continue
            if cand.date_source == "first archive capture" and cand.date and cand.date < start - dt.timedelta(days=60):
                stats["skipped_out_of_range"] += 1
                continue
            if (src.title_prefilter or src.slug_prefilter) and cand.title_hint \
                    and not match_keywords(cand.title_hint, keywords)["distinct"]:
                stats["skipped_title_prefilter"] += 1
                continue
            if limit and stats["fetched_articles"] >= limit:
                break
            page = fetcher.get(cand.fetch_url or cand.url, render_js=src.render_js, when=cand.date)
            if not page:
                reason = fetcher.last_reason.value or "failed"
                stats["blocked" if reason in ("blocked", "http_403") else f"fetch_{reason}"] += 1
                continue
            stats["fetched_articles"] += 1
            if page.via.endswith("wayback"):
                stats["via_wayback"] += 1
            meta = extract(page, cand.url)
            title = (meta.get("title") or cand.title_hint or "").strip()
            if cand.title_hint and src.strategy in ("federal_register", "govuk_api"):
                title = cand.title_hint
            text = (meta.get("text") or "").strip()
            date, date_src = resolve_date(cand, meta)
            if date and not (start <= date <= end):
                stats["dropped_out_of_range"] += 1
                continue
            if len(text.split()) < 30:
                stats["dropped_empty_text"] += 1
                continue
            m = match_keywords(title + "\n" + text, keywords)
            if not m["distinct"]:
                stats["dropped_no_keywords"] += 1
                continue
            stats["kept"] += 1
            north_south, continent = COUNTRY_INFO.get(src.country, ("", ""))
            records.append({
                "month": f"{date:%Y-%m}" if date else "Undated", "date": date, "period": period_of(date),
                "country": src.country, "north_south": north_south, "continent": continent,
                "source_type": src.source_type,
                "doc_type": classify_doc_type(title, cand.url, src.source_type, cand.doc_type),
                "site": urlparse(cand.url).netloc.removeprefix("www."), "title": title, "url": cand.url,
                "matched_keywords": "; ".join(sorted(set().union(*m["by_cat"].values()))),
                "by_cat": {k: sorted(v) for k, v in m["by_cat"].items()},
                "keyword_count": m["distinct"], "strong_match": "Yes" if m["strong"] else "No",
                "person_countries": ", ".join(m["person_countries"]),
                "date_source": date_src, "fetched_via": page.via.removeprefix("cache:"),
                "word_count": len(text.split()), "text": text,
            })
    except Exception as e:
        log.exception("source failed: %s", src.url)
        stats["error_msg"] = str(e)[:200]
    log.info("  %-58s candidates=%-5d fetched=%-5d kept=%-4d blocked=%d", src.url[:58], stats["candidates"],
             stats["fetched_articles"], stats["kept"], stats["blocked"])
    return records, stats


# ─────────────────────────────────────────────────────────────────────────────
# 9. PROBE: quick health check of one country's sources
# ─────────────────────────────────────────────────────────────────────────────
def probe(sources, fetcher, start, end, out_path: Path):
    rows = []
    for src in sources:
        row = {"country": src.country, "type": src.source_type, "url": src.url, "strategy": src.strategy,
               "enabled": src.enabled, "notes": src.notes}
        if not src.enabled:
            row.update(status="disabled", links_found=0)
        elif src.strategy in ("govuk_api", "hk_daily", "federal_register", "wayback", "manual"):
            if src.strategy == "wayback":
                gen = discover_wayback(src, fetcher, start, end, 1, cdx_limit=50)
            elif src.strategy == "hk_daily":
                gen = discover_hk_daily(src, fetcher, start, min(end, start + dt.timedelta(days=5)), 1)
            else:
                gen = STRATEGIES[src.strategy](src, fetcher, start, end, 1)
            sample = [c for _, c in zip(range(5), gen)]
            row.update(status="ok" if sample else f"no results ({fetcher.last_reason.value or 'empty'})",
                       links_found=len(sample), sample_links=" | ".join(c.url for c in sample[:3]))
        else:
            first = page_url(src, 0, start, end)
            page = fetcher.get(first, max_age=DAY, render_js=src.render_js, allow_wayback=False)
            if not page:
                reason = fetcher.last_reason.value
                hint = {"blocked": "BLOCKED: run on your own computer / --js; Wayback fallback will be used",
                        "needs_js": "needs --js (JavaScript site)", "robots": "disallowed by robots.txt"}
                row.update(status=hint.get(reason, f"fetch failed ({reason})"), links_found=0,
                           wayback_fallback="yes" if src.wayback_prefix else "no (add wayback_prefix)")
            else:
                links = extract_links(page.text, first, src)
                row.update(status="ok" if links else "0 article links: check link_pattern / JS",
                           links_found=len(links), dated_links=sum(bool(c.date) for c in links),
                           has_next_link=bool(find_next_link(page.text, first)),
                           sample_links=" | ".join(c.url for c in links[:3]))
        rows.append(row)
        log.info("  %-62s %s (%s links)", src.url[:62], row["status"], row.get("links_found"))
    pd.DataFrame(rows).to_csv(out_path, index=False)
    log.info("Probe report: %s", out_path)


# ─────────────────────────────────────────────────────────────────────────────
# 10. OUTPUTS
# ─────────────────────────────────────────────────────────────────────────────
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="2F5597")
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
RED_FILL = PatternFill("solid", start_color="F8CBAD")

COLS = [  # header, key, width   (Month must stay col A, Period C, Country D, Strong Match M)
    ("Month", "month", 10), ("Date", "date", 12), ("Period", "period", 15), ("Country", "country", 14),
    ("North/South", "north_south", 13), ("Continent", "continent", 13), ("Source Type", "source_type", 22),
    ("Document Type", "doc_type", 22), ("Site", "site", 24), ("Title", "title", 60), ("URL", "url", 48),
    ("Matched Keywords", "matched_keywords", 42), ("Strong Match", "strong_match", 10),
    ("Keyword Count", "keyword_count", 10), ("Minister/Person Country", "person_countries", 16),
    ("Date Source", "date_source", 16), ("Fetched Via", "fetched_via", 11), ("Word Count", "word_count", 10),
]
COL_LETTER = {k: get_column_letter(i) for i, (_, k, _) in enumerate(COLS, start=1)}
LOG_COLS = ["country", "source_type", "url", "strategy", "enabled", "candidates", "fetched_articles", "kept",
            "blocked", "via_wayback", "skipped_out_of_range", "skipped_title_prefilter", "dropped_out_of_range",
            "dropped_empty_text", "dropped_no_keywords", "fetch_needs_js", "fetch_failed", "error_msg", "notes"]


def _clean(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _header(ws, headers, widths):
    ws.append(headers)
    for i, (c, w) in enumerate(zip(ws[1], widths), start=1):
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _style_row(ws, font=BODY):
    for c in ws[ws.max_row]:
        c.font = font


def months_between(start, end):
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        yield f"{y}-{m:02d}"
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def _article_sheet(wb, name, records, cats):
    ws = wb.create_sheet(name)
    _header(ws, [h for h, _, _ in COLS] + [f"{c} hits" for c in cats], [w for *_, w in COLS] + [26] * len(cats))
    url_col = [k for _, k, _ in COLS].index("url") + 1
    for r in records:
        ws.append([_clean(r.get(k, "")) for _, k, _ in COLS] + ["; ".join(r["by_cat"].get(c, [])) for c in cats])
        _style_row(ws)
        cell = ws.cell(ws.max_row, url_col)
        cell.hyperlink, cell.font = r["url"], LINK
        ws.cell(ws.max_row, 2).number_format = "yyyy-mm-dd"
    if records:
        ws.auto_filter.ref = ws.dimensions


def write_workbook(path: Path, records, source_rows, countries, start, end, title_note=""):
    records = sorted(records, key=lambda r: (r["date"] or dt.date.max, r["country"], r["title"]))
    cats = list(dict.fromkeys(c for r in records for c in r["by_cat"]))
    wb = Workbook()
    wb.remove(wb.active)
    A, Cn, Pc, Sm = COL_LETTER["month"], COL_LETTER["country"], COL_LETTER["period"], COL_LETTER["strong_match"]
    NS = COL_LETTER["north_south"]
    AA = "'All Articles'"

    # Monthly Summary (live formulas over All Articles)
    ws = wb.create_sheet("Monthly Summary")
    _header(ws, ["Month", "Period"] + countries + ["Total", "Strong matches", "Global North", "Global South"],
            [10, 15] + [12] * len(countries) + [9, 12, 12, 12])
    last_c = get_column_letter(2 + len(countries))
    first_row = 2
    for i, month in enumerate(months_between(start, end), start=first_row):
        row = [month, period_of(dt.date.fromisoformat(month + "-01"))]
        row += [f"=COUNTIFS({AA}!${A}:${A},$A{i},{AA}!${Cn}:${Cn},{get_column_letter(3 + j)}$1)"
                for j in range(len(countries))]
        row += [f"=SUM(C{i}:{last_c}{i})", f"=COUNTIFS({AA}!${A}:${A},$A{i},{AA}!${Sm}:${Sm},\"Yes\")",
                f"=COUNTIFS({AA}!${A}:${A},$A{i},{AA}!${NS}:${NS},\"Global North\")",
                f"=COUNTIFS({AA}!${A}:${A},$A{i},{AA}!${NS}:${NS},\"Global South\")"]
        ws.append(row)
        _style_row(ws)
    last = ws.max_row
    ws.append(["Total", ""] + [f"=SUM({get_column_letter(c)}{first_row}:{get_column_letter(c)}{last})"
                               for c in range(3, 7 + len(countries))])
    _style_row(ws, BOLD)
    # months with zero articles for a country are shaded: these are the coverage gaps
    ws.conditional_formatting.add(f"C{first_row}:{last_c}{last}",
                                  CellIsRule(operator="equal", formula=["0"], fill=RED_FILL))
    ws.append([])
    ws.append(["Undated articles", "", f"=COUNTIFS({AA}!${A}:${A},\"Undated\")"])
    _style_row(ws)
    ws.append([])
    ws.append(["Period", ""] + countries + ["Total"])
    hdr = ws.max_row
    _style_row(ws, BOLD)
    for label, _, _ in PERIODS:
        r = ws.max_row + 1
        ws.append([label, ""] + [f"=COUNTIFS({AA}!${Pc}:${Pc},$A{r},{AA}!${Cn}:${Cn},{get_column_letter(3 + j)}${hdr})"
                                 for j in range(len(countries))] + [f"=SUM(C{r}:{last_c}{r})"])
        _style_row(ws)
    ws.append([])
    ws.append(["Shaded cells = no articles collected for that country in that month (a coverage gap)."])
    ws[ws.max_row][0].font = Font(name="Arial", size=9, italic=True)

    _article_sheet(wb, "All Articles", records, cats)
    by_month = defaultdict(list)
    for r in records:
        by_month[r["month"]].append(r)
    for month in sorted(by_month, key=lambda m: (m == "Undated", m)):
        _article_sheet(wb, month, by_month[month], cats)

    ws = wb.create_sheet("Source Log")
    _header(ws, [c.replace("_", " ").title() for c in LOG_COLS], [13, 20, 50, 12, 8] + [11] * 12 + [36, 55])
    for row in source_rows:
        ws.append([_clean(row.get(c, "")) for c in LOG_COLS])
        _style_row(ws)

    ws = wb.create_sheet("Notes")
    _header(ws, ["Item", "Detail"], [26, 110])
    for k, v in [
        ("Built", dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("Date range", f"{start} to {end}"),
        ("Scope", title_note),
        ("Periods (comment NJ1)", "; ".join(f"{l}: {a} to {b} (excl.)" for l, a, b in PERIODS)),
        ("North/South", "Grouping from the Step 1 text; continent from the sheet's note."),
        ("Document Type", "Rule-based from title/URL (executive order, speech, statement, policy, press release)."),
        ("Strong Match", f"Any non-indicator codeword, or {MIN_BROAD_TERMS_FOR_STRONG}+ distinct indicator words."),
        ("Date Source", "url / API / daily index > page metadata > listing page > first archive capture (approximate)."),
        ("Fetched Via", "live = the site; js = real browser; wayback = Internet Archive copy (site blocked us)."),
        ("Blocked", "Source Log 'Blocked' = pages the site refused even after fallbacks. See README."),
    ]:
        ws.append([k, v])
        _style_row(ws)
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def save_country(records, source_rows, country, out_dir: Path, start, end):
    out_dir.mkdir(parents=True, exist_ok=True)
    records.sort(key=lambda r: (r["date"] or dt.date.max, r["title"]))
    with open(out_dir / "corpus.jsonl", "w", encoding="utf-8") as f:
        for r in records:
            f.write(json.dumps({**{k: v for k, v in r.items() if k != "by_cat"},
                                "date": r["date"].isoformat() if r["date"] else None,
                                "keywords_by_category": r["by_cat"]}, ensure_ascii=False) + "\n")
    pd.DataFrame([{h: r.get(k, "") for h, k, _ in COLS} for r in records], columns=[h for h, _, _ in COLS]) \
        .to_csv(out_dir / "articles.csv", index=False)
    pd.DataFrame(source_rows, columns=LOG_COLS).to_csv(out_dir / "source_log.csv", index=False)
    (out_dir / "run_info.json").write_text(json.dumps(
        {"country": country, "start": str(start), "end": str(end), "articles": len(records),
         "finished": dt.datetime.now().isoformat(timespec="seconds")}, indent=2))
    return write_workbook(out_dir / f"{slug(country)}_Articles.xlsx", records, source_rows, [country],
                          start, end, f"Country: {country}")


def combine(out_root: Path, start, end):
    """Merge outputs/countries/*/corpus.jsonl into one month-sorted workbook."""
    folders = sorted(p for p in (out_root / "countries").glob("*") if (p / "corpus.jsonl").exists())
    if not folders:
        raise SystemExit(f"No finished countries in {out_root / 'countries'}. Run --country first.")
    records, logs, seen, dups = [], [], set(), 0
    for f in folders:
        for line in open(f / "corpus.jsonl", encoding="utf-8"):
            r = json.loads(line)
            if r["url"] in seen:
                dups += 1
                continue
            seen.add(r["url"])
            r["date"] = iso_date(r.get("date"))
            r["by_cat"] = r.pop("keywords_by_category", {})
            if r["date"] and not (start <= r["date"] <= end):
                continue
            records.append(r)
        if (f / "source_log.csv").exists():
            logs += pd.read_csv(f / "source_log.csv", dtype=str).fillna("").to_dict("records")
    order = list(COUNTRY_INFO)
    countries = sorted({r["country"] for r in records} | {json.loads((f / "run_info.json").read_text())["country"]
                                                          for f in folders if (f / "run_info.json").exists()},
                       key=lambda c: (order.index(c) if c in order else 99, c))
    path = write_workbook(out_root / "Caraway_Jain_Articles_ALL.xlsx", records, logs, countries, start, end,
                          f"Combined: {', '.join(countries)}")
    with open(out_root / "corpus_all.jsonl", "w", encoding="utf-8") as fh:
        for r in sorted(records, key=lambda r: (r["date"] or dt.date.max, r["country"])):
            fh.write(json.dumps({**{k: v for k, v in r.items() if k != "by_cat"},
                                 "date": r["date"].isoformat() if r["date"] else None,
                                 "keywords_by_category": r["by_cat"]}, ensure_ascii=False) + "\n")
    pd.DataFrame([{h: r.get(k, "") for h, k, _ in COLS} for r in records]).to_csv(out_root / "articles_all.csv", index=False)
    missing = [c for c in COUNTRY_INFO if c not in countries]
    log.info("Combined %d articles from %d countries (%d cross-country duplicates removed) -> %s",
             len(records), len(folders), dups, path)
    if missing:
        log.info("Not yet run: %s", ", ".join(missing))
    return path


def slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")


# ─────────────────────────────────────────────────────────────────────────────
# 11. MAIN
# ─────────────────────────────────────────────────────────────────────────────
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--country", help="run ONE country, e.g. USA, 'Hong Kong', 'International (news)'")
    ap.add_argument("--combine", action="store_true", help="merge all finished countries into one workbook")
    ap.add_argument("--list", action="store_true", help="list countries and their sources")
    ap.add_argument("--probe", action="store_true", help="health-check the country's sources only")
    ap.add_argument("--excel", default=None, help="sheet .xlsx, or folder with General/Codewords CSVs")
    ap.add_argument("--start", default=DEFAULT_START)
    ap.add_argument("--end", default=DEFAULT_END)
    ap.add_argument("--out", default="outputs")
    ap.add_argument("--manual-dir", default="manual_urls")
    ap.add_argument("--types", default="gov,ministry,other", help="gov,ministry,news,other (other = API/manual)")
    ap.add_argument("--max-pages", type=int, default=400, help="max listing pages per source")
    ap.add_argument("--limit-per-source", type=int, default=0, help="trial runs only (0 = no limit)")
    ap.add_argument("--delay", type=float, default=2.0, help="seconds between hits on one site")
    ap.add_argument("--workers", type=int, default=4, help="sources of this country run in parallel")
    ap.add_argument("--email", default=CONTACT_EMAIL)
    ap.add_argument("--js", action="store_true", help="use a real browser (Playwright) for JS/blocked sites")
    ap.add_argument("--browser", default="chromium", choices=["chromium", "chrome", "msedge"],
                    help="with --js: 'chrome' uses your installed Google Chrome")
    ap.add_argument("--headed", action="store_true", help="with --js: show the browser window")
    ap.add_argument("--no-wayback", action="store_true", help="never use Internet Archive copies")
    ap.add_argument("--offline", action="store_true", help="cached pages only")
    ap.add_argument("-v", "--verbose", action="store_true")
    args, _ = ap.parse_known_args(argv)

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S", force=True)
    for noisy in ("trafilatura", "htmldate", "urllib3", "charset_normalizer", "pypdf"):
        logging.getLogger(noisy).setLevel(logging.CRITICAL)

    start, end = dt.date.fromisoformat(args.start), dt.date.fromisoformat(args.end)
    out_root = Path(args.out)
    if args.combine:
        return combine(out_root, start, end)

    sheet = find_sheet(args.excel)
    sources = load_sources(sheet, Path(args.manual_dir))
    keywords = load_keywords(sheet)
    countries = list(dict.fromkeys(s.country for s in sources))
    if args.list or not args.country:
        print(f"Sheet: {sheet.resolve()}  |  {len(keywords)} keywords")
        for c in countries:
            print(f"\n{c}")
            for s in sources:
                if s.country == c:
                    print(f"   [{'on ' if s.enabled else 'OFF'}] {s.strategy:16} {s.source_type[:24]:24} {s.url[:80]}")
        if not args.country:
            print('\nNext: main(["--country", "<one of the names above>"])')
        return

    match = [c for c in countries if c.lower() == args.country.strip().lower()]
    if not match:
        raise SystemExit(f"Unknown country '{args.country}'. Choose from: {', '.join(countries)}")
    country = match[0]
    type_map = {"gov": SOURCE_TYPE_LABELS["gov"], "ministry": SOURCE_TYPE_LABELS["ministry"],
                "news": SOURCE_TYPE_LABELS["news"]}
    wanted = [t.strip() for t in args.types.split(",")]
    def type_ok(s):
        known = s.source_type in type_map.values()
        return (known and any(type_map.get(t) == s.source_type for t in wanted)) or (not known and "other" in wanted) \
            or country == "International (news)"
    srcs = [s for s in sources if s.country == country and type_ok(s)]
    out_dir = out_root / "countries" / slug(country)
    out_dir.mkdir(parents=True, exist_ok=True)
    log.info("%s: %d sources, %d keywords, %s to %s", country, len(srcs), len(keywords), start, end)

    fetcher = Fetcher(Path("cache") / f"{slug(country)}.sqlite", args.delay, offline=args.offline, use_js=args.js,
                      browser=args.browser, headed=args.headed, wayback=not args.no_wayback, email=args.email)
    try:
        if args.probe:
            return probe(srcs, fetcher, start, end, out_dir / "probe_report.csv")
        records, source_rows = [], []
        active = [s for s in srcs if s.enabled]
        source_rows += [{**vars(s), "error_msg": "disabled"} for s in srcs if not s.enabled]
        workers = 1 if args.js else max(1, args.workers)
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futs = {pool.submit(run_source, s, fetcher, keywords, start, end, args.max_pages,
                                args.limit_per_source): s for s in active}
            for fut in as_completed(futs):
                s = futs[fut]
                recs, stats = fut.result()
                records += recs
                source_rows.append({**{k: v for k, v in vars(s).items() if k != "manual_rows"}, **stats})
        dedup, seen = [], set()
        for r in records:
            if r["url"] not in seen:
                seen.add(r["url"]); dedup.append(r)
        path = save_country(dedup, source_rows, country, out_dir, start, end)
        blocked = sum(int(r.get("blocked", 0) or 0) for r in source_rows)
        log.info("%s done: %d articles -> %s", country, len(dedup), path)
        log.info("Fetch summary: %s", dict(fetcher.stats))
        if blocked:
            log.warning("%d pages were blocked even after fallbacks. Busiest blocked hosts: %s. See README section 'Blocking'.",
                        blocked, dict(fetcher.blocked_hosts.most_common(4)))
        return path
    finally:
        fetcher.close()


if __name__ == "__main__":
    main()
